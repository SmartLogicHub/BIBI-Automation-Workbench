from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


HEADERS = ["video_url", "account_name", "comment_text", "status", "error_reason", "created_at", "updated_at"]
SUCCESS_STATUS = "success"
FAILURE_STATUSES = {
    "failed",
    "manual_timeout",
    "login_required",
    "captcha_or_verify",
    "comment_closed",
    "submit_failed",
    "page_error",
}


@dataclass(frozen=True)
class Record:
    video_url: str
    account_name: str
    comment_text: str
    status: str
    error_reason: str
    created_at: str
    updated_at: str


class RecordStorage:
    def __init__(self, posted_file, failed_file, processed_file):
        self.posted_file = Path(posted_file)
        self.failed_file = Path(failed_file)
        self.processed_file = Path(processed_file)
        for path in [self.posted_file, self.failed_file, self.processed_file]:
            self._ensure_workbook(path)

    def _ensure_workbook(self, path):
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.exists():
            return
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(HEADERS)
        workbook.save(path)

    def _append(self, path, record):
        workbook = load_workbook(path)
        try:
            sheet = workbook.active
            sheet.append([
                record.video_url,
                record.account_name,
                record.comment_text,
                record.status,
                record.error_reason,
                record.created_at,
                record.updated_at,
            ])
            workbook.save(path)
        finally:
            workbook.close()

    def record(self, video_url, account_name, comment_text, status, error_reason=""):
        now = datetime.now().astimezone().isoformat(timespec="seconds")
        record = Record(
            video_url=str(video_url or "").strip(),
            account_name=str(account_name or "").strip(),
            comment_text=str(comment_text or ""),
            status=str(status or "").strip(),
            error_reason=str(error_reason or ""),
            created_at=now,
            updated_at=now,
        )
        self._append(self.processed_file, record)
        if record.status == SUCCESS_STATUS:
            self._append(self.posted_file, record)
        elif record.status in FAILURE_STATUSES:
            self._append(self.failed_file, record)
        return record

    def _records_from(self, path):
        if not Path(path).exists():
            return []
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        if len(rows) <= 1:
            return []
        headers = {str(cell or "").strip(): index for index, cell in enumerate(rows[0])}
        records = []
        for row in rows[1:]:
            records.append(
                Record(
                    video_url=str(row[headers["video_url"]] or "").strip(),
                    account_name=str(row[headers["account_name"]] or "").strip(),
                    comment_text=str(row[headers["comment_text"]] or ""),
                    status=str(row[headers["status"]] or "").strip(),
                    error_reason=str(row[headers["error_reason"]] or ""),
                    created_at=str(row[headers["created_at"]] or ""),
                    updated_at=str(row[headers["updated_at"]] or ""),
                )
            )
        return records

    def processed_records(self):
        return self._records_from(self.processed_file)

    def success_urls(self):
        return {record.video_url for record in self.processed_records() if record.status == SUCCESS_STATUS}
