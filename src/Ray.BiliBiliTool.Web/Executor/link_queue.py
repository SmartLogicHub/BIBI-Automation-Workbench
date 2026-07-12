from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class LinkItem:
    video_url: str
    title: str = ""
    source_keyword: str = ""
    created_at: str = ""


def _normalize_url(value):
    return str(value or "").strip()


def _header_map(row):
    return {str(cell or "").strip(): index for index, cell in enumerate(row)}


def build_link_queue(input_file, success_urls=None):
    path = Path(input_file)
    if not path.exists():
        raise FileNotFoundError(f"链接文件不存在：{path}")

    success_urls = {_normalize_url(url) for url in (success_urls or set()) if _normalize_url(url)}
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []

    headers = _header_map(rows[0])
    if "video_url" not in headers:
        raise ValueError("链接文件必须包含 video_url 字段")

    seen = set()
    queue = []
    for row in rows[1:]:
        video_url = _normalize_url(row[headers["video_url"]] if headers["video_url"] < len(row) else "")
        if not video_url or video_url in seen or video_url in success_urls:
            continue
        seen.add(video_url)
        queue.append(
            LinkItem(
                video_url=video_url,
                title=str(row[headers["title"]] or "").strip() if "title" in headers and headers["title"] < len(row) else "",
                source_keyword=str(row[headers["source_keyword"]] or "").strip()
                if "source_keyword" in headers and headers["source_keyword"] < len(row)
                else "",
                created_at=str(row[headers["created_at"]] or "").strip()
                if "created_at" in headers and headers["created_at"] < len(row)
                else "",
            )
        )
    return queue
